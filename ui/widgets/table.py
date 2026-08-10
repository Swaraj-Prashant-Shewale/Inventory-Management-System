"""DataTable — the table used on every list screen.

Handles column sizing, numeric-aware sorting, an optional tick column with a working
select-all header, an empty-state message, and CSV/Excel export.
"""
import csv
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Optional

from PySide6.QtCore import QPoint, QRect, Qt, Signal
from PySide6.QtGui import QBrush, QColor, QKeySequence, QPainter, QPen
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QHeaderView,
    QLabel,
    QTableWidget,
    QTableWidgetItem,
)

from ui import theme

LEFT = Qt.AlignLeft | Qt.AlignVCenter
CENTER = Qt.AlignCenter
RIGHT = Qt.AlignRight | Qt.AlignVCenter


@dataclass
class Column:
    """One table column. `width` is 'stretch', 'content', or a pixel integer."""
    title: str
    width: Any = "stretch"
    align: Any = LEFT
    tooltip: Optional[str] = None


@dataclass
class Cell:
    """One table cell. `sort_value` keeps numeric columns sorting numerically."""
    text: str = ""
    align: Any = None
    sort_value: Any = None
    colour: Optional[str] = None
    bold: bool = False
    tooltip: Optional[str] = None


@dataclass
class Row:
    cells: list
    payload: Any = None
    tint: Optional[str] = None       # row background, e.g. to flag low stock
    checkable: bool = True
    meta: dict = field(default_factory=dict)


class _SortableItem(QTableWidgetItem):
    """Sorts on `sort_value` when present, so ₹1,00,000 beats ₹9,999 correctly."""

    def __init__(self, text=""):
        super().__init__(text)
        self.sort_value = None

    def __lt__(self, other):
        mine = self.sort_value
        theirs = getattr(other, "sort_value", None)
        if mine is None and theirs is None:
            # Never delegate to QTableWidgetItem.__lt__ here: Qt routes it straight back
            # into this override, which recurses until the stack blows.
            return self.text().casefold() < other.text().casefold()
        if mine is None:
            return True
        if theirs is None:
            return False
        try:
            return mine < theirs
        except TypeError:
            return str(mine) < str(theirs)


class _CheckHeader(QHeaderView):
    """Header with a working select-all tick in the first section."""

    toggled = Signal(bool)

    def __init__(self, parent=None):
        super().__init__(Qt.Horizontal, parent)
        self.setSectionsClickable(True)
        self._checked = False

    def set_checked(self, checked, emit=False):
        if self._checked != checked:
            self._checked = checked
            self.viewport().update()
            if emit:
                self.toggled.emit(checked)

    def is_checked(self):
        return self._checked

    def paintSection(self, painter, rect, index):
        painter.save()
        super().paintSection(painter, rect, index)
        painter.restore()
        if index != 0:
            return

        # Drawn by hand rather than via CE_CheckBox so it matches the rounded, teal
        # checkboxes the stylesheet gives every other checkbox in the app.
        size = 19
        box = QRect(
            rect.x() + (rect.width() - size) // 2,
            rect.y() + (rect.height() - size) // 2,
            size, size,
        )
        painter.save()
        painter.setRenderHint(QPainter.Antialiasing)
        if self._checked:
            painter.setBrush(QColor(theme.TEAL))
            painter.setPen(QPen(QColor(theme.TEAL), 2))
        else:
            painter.setBrush(QColor("#FFFFFF"))
            painter.setPen(QPen(QColor(theme.BORDER_STRONG), 2))
        painter.drawRoundedRect(box.adjusted(1, 1, -1, -1), 4, 4)

        if self._checked:
            tick = QPen(QColor("#FFFFFF"), 2.2)
            tick.setCapStyle(Qt.RoundCap)
            tick.setJoinStyle(Qt.RoundJoin)
            painter.setPen(tick)
            painter.drawPolyline([
                QPoint(box.left() + 5, box.center().y()),
                QPoint(box.center().x() - 1, box.bottom() - 5),
                QPoint(box.right() - 4, box.top() + 5),
            ])
        painter.restore()

    def mousePressEvent(self, event):
        if self.logicalIndexAt(event.position().toPoint()) == 0:
            self.set_checked(not self._checked, emit=True)
            return
        super().mousePressEvent(event)


class DataTable(QTableWidget):
    """List table with optional tick column."""

    selection_changed = Signal()
    check_changed = Signal()
    row_activated = Signal(object)        # emits the payload on double-click

    def __init__(self, columns, checkable=False, parent=None,
                 empty_text="Nothing to show yet."):
        super().__init__(parent)
        self._columns = list(columns)
        self._checkable = checkable
        self._rows = []
        self._offset = 1 if checkable else 0

        headers = ([""] if checkable else []) + [c.title for c in self._columns]
        self.setColumnCount(len(headers))
        self.setHorizontalHeaderLabels(headers)

        if checkable:
            header = _CheckHeader(self)
            self.setHorizontalHeader(header)
            self.setHorizontalHeaderLabels(headers)
            header.toggled.connect(self.set_all_checked)

        self.verticalHeader().setVisible(False)
        self.verticalHeader().setDefaultSectionSize(44)
        self.setShowGrid(False)
        self.setAlternatingRowColors(False)
        self.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.setWordWrap(False)
        self.setSortingEnabled(True)
        self.setFocusPolicy(Qt.StrongFocus)
        self.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)

        self._apply_column_sizing()

        self.itemSelectionChanged.connect(self.selection_changed.emit)
        self.itemChanged.connect(self._on_item_changed)
        self.itemDoubleClicked.connect(self._on_double_click)

        self._empty_label = QLabel(empty_text, self.viewport())
        self._empty_label.setAlignment(Qt.AlignCenter)
        self._empty_label.setStyleSheet(
            f"color: {theme.TEXT_MUTED}; font-size: 14px; background: transparent;"
        )
        self._empty_label.hide()

    # --- Layout ---------------------------------------------------------------------

    def _apply_column_sizing(self):
        header = self.horizontalHeader()
        header.setHighlightSections(False)
        # No sort arrow until the user actually sorts — otherwise one appears over the
        # select-all tick in column 0 on first paint.
        header.setSortIndicator(-1, Qt.AscendingOrder)
        if self._checkable:
            header.setSectionResizeMode(0, QHeaderView.Fixed)
            self.setColumnWidth(0, 48)
        for i, column in enumerate(self._columns):
            index = i + self._offset
            if column.width == "stretch":
                header.setSectionResizeMode(index, QHeaderView.Stretch)
            elif column.width == "content":
                header.setSectionResizeMode(index, QHeaderView.ResizeToContents)
            else:
                header.setSectionResizeMode(index, QHeaderView.Fixed)
                self.setColumnWidth(index, int(column.width))
            if column.tooltip:
                item = self.horizontalHeaderItem(index)
                if item:
                    item.setToolTip(column.tooltip)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._empty_label.setGeometry(self.viewport().rect())

    # --- Data -----------------------------------------------------------------------

    def set_rows(self, rows):
        """Replace all rows. `rows` is a list of Row objects."""
        checked_before = {id(p) for p in self.checked_payloads()}
        self.setSortingEnabled(False)
        self.blockSignals(True)
        self.clearContents()

        self._rows = list(rows)
        self.setRowCount(len(self._rows))

        for r, row in enumerate(self._rows):
            if self._checkable:
                tick = _SortableItem("")
                if row.checkable:
                    tick.setFlags(
                        Qt.ItemIsUserCheckable | Qt.ItemIsEnabled | Qt.ItemIsSelectable
                    )
                    tick.setCheckState(
                        Qt.Checked if id(row.payload) in checked_before else Qt.Unchecked
                    )
                else:
                    tick.setFlags(Qt.ItemIsEnabled)
                tick.setTextAlignment(CENTER)
                tick.setData(Qt.UserRole, row.payload)
                self.setItem(r, 0, tick)

            for c, column in enumerate(self._columns):
                value = row.cells[c] if c < len(row.cells) else ""
                cell = value if isinstance(value, Cell) else Cell(text=str(value))
                item = _SortableItem(cell.text or "")
                item.setTextAlignment(cell.align if cell.align is not None else column.align)
                item.sort_value = (
                    cell.sort_value if cell.sort_value is not None
                    else _auto_sort_value(cell.text)
                )
                if cell.colour:
                    item.setForeground(QBrush(QColor(cell.colour)))
                if cell.bold:
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                if cell.tooltip:
                    item.setToolTip(cell.tooltip)
                if row.tint:
                    item.setBackground(QBrush(QColor(row.tint)))
                item.setData(Qt.UserRole, row.payload)
                self.setItem(r, c + self._offset, item)

        self.blockSignals(False)
        self.setSortingEnabled(True)
        self._empty_label.setVisible(not self._rows)
        self._empty_label.raise_()
        self._sync_header_check()
        self.selection_changed.emit()
        self.check_changed.emit()

    def set_empty_text(self, text):
        self._empty_label.setText(text)

    def row_count(self):
        return len(self._rows)

    # --- Selection ------------------------------------------------------------------

    def current_payload(self):
        row = self.currentRow()
        if row < 0:
            return None
        item = self.item(row, self._offset)
        return item.data(Qt.UserRole) if item else None

    def selected_payloads(self):
        payloads, seen = [], set()
        for index in self.selectedIndexes():
            item = self.item(index.row(), self._offset)
            if item is None:
                continue
            payload = item.data(Qt.UserRole)
            if payload is not None and id(payload) not in seen:
                seen.add(id(payload))
                payloads.append(payload)
        return payloads

    def checked_payloads(self):
        if not self._checkable:
            return []
        result = []
        for r in range(self.rowCount()):
            item = self.item(r, 0)
            if item is not None and item.checkState() == Qt.Checked:
                payload = item.data(Qt.UserRole)
                if payload is not None:
                    result.append(payload)
        return result

    def target_payloads(self):
        """Ticked rows if any, otherwise the highlighted rows.

        Lets a user either tick several rows or just click one and press the button.
        """
        return self.checked_payloads() or self.selected_payloads()

    def set_all_checked(self, checked):
        if not self._checkable:
            return
        self.blockSignals(True)
        state = Qt.Checked if checked else Qt.Unchecked
        for r in range(self.rowCount()):
            item = self.item(r, 0)
            if item is not None and item.flags() & Qt.ItemIsUserCheckable:
                item.setCheckState(state)
        self.blockSignals(False)
        self.check_changed.emit()

    def clear_checks(self):
        self.set_all_checked(False)
        header = self.horizontalHeader()
        if isinstance(header, _CheckHeader):
            header.set_checked(False)

    def _on_item_changed(self, item):
        if self._checkable and item.column() == 0:
            self._sync_header_check()
            self.check_changed.emit()

    def _sync_header_check(self):
        header = self.horizontalHeader()
        if not isinstance(header, _CheckHeader):
            return
        total = sum(
            1 for r in range(self.rowCount())
            if self.item(r, 0) is not None
            and self.item(r, 0).flags() & Qt.ItemIsUserCheckable
        )
        checked = len(self.checked_payloads())
        header.set_checked(total > 0 and checked == total)

    def _on_double_click(self, item):
        payload = item.data(Qt.UserRole)
        if payload is not None:
            self.row_activated.emit(payload)

    # --- Export ---------------------------------------------------------------------

    def visible_headers(self):
        return [c.title for c in self._columns]

    def visible_matrix(self):
        """Rows as plain strings, in the order currently displayed on screen."""
        data = []
        for r in range(self.rowCount()):
            if self.isRowHidden(r):
                continue
            data.append([
                (self.item(r, c + self._offset).text()
                 if self.item(r, c + self._offset) else "")
                for c in range(len(self._columns))
            ])
        return data

    def export_csv(self, path, title=None):
        with open(path, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            if title:
                writer.writerow([title])
                writer.writerow([])
            writer.writerow(self.visible_headers())
            writer.writerows(self.visible_matrix())
        return path

    def export_xlsx(self, path, title="Export"):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        book = Workbook()
        sheet = book.active
        sheet.title = title[:31] or "Export"

        headers = self.visible_headers()
        sheet.append(headers)
        fill = PatternFill("solid", fgColor="005F6B")
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(row=1, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in self.visible_matrix():
            sheet.append(row)

        for col in range(1, len(headers) + 1):
            longest = max(
                [len(str(headers[col - 1]))]
                + [len(str(sheet.cell(row=r, column=col).value or ""))
                   for r in range(2, sheet.max_row + 1)]
            )
            sheet.column_dimensions[get_column_letter(col)].width = min(48, longest + 4)

        sheet.freeze_panes = "A2"
        book.save(path)
        return path

    def copy_selection(self):
        """Copy highlighted rows to the clipboard as TSV, for pasting into Excel."""
        rows = sorted({i.row() for i in self.selectedIndexes()})
        if not rows:
            return
        lines = []
        for r in rows:
            lines.append("\t".join(
                (self.item(r, c + self._offset).text()
                 if self.item(r, c + self._offset) else "")
                for c in range(len(self._columns))
            ))
        QApplication.clipboard().setText("\n".join(lines))

    def keyPressEvent(self, event):
        if event.matches(QKeySequence.Copy):
            self.copy_selection()
            return
        super().keyPressEvent(event)


def _auto_sort_value(text):
    """Best-effort numeric sort key so money and quantity columns order correctly."""
    if not text:
        return None
    cleaned = str(text).strip().replace(",", "").replace("₹", "").replace("%", "")
    cleaned = cleaned.replace(" hrs", "").strip()
    if not cleaned or cleaned in {"—", "-"}:
        return None
    try:
        return Decimal(cleaned)
    except (InvalidOperation, ValueError):
        return None


def money_cell(value, symbol="₹", tone=None, bold=False):
    from decimal import Decimal as D
    try:
        raw = D(str(value or 0))
    except (InvalidOperation, ValueError):
        raw = D(0)
    return Cell(theme.money(raw, symbol), align=RIGHT, sort_value=raw,
                colour=tone, bold=bold)


def qty_cell(value, tone=None, bold=False, suffix=""):
    from decimal import Decimal as D
    try:
        raw = D(str(value or 0))
    except (InvalidOperation, ValueError):
        raw = D(0)
    return Cell(theme.quantity(raw) + suffix, align=RIGHT, sort_value=raw,
                colour=tone, bold=bold)


def text_cell(value, align=LEFT, tone=None, bold=False, tooltip=None):
    return Cell("" if value is None else str(value), align=align, colour=tone,
                bold=bold, tooltip=tooltip)


def date_cell(value, align=CENTER):
    return Cell(theme.date_short(value), align=align, sort_value=value)
